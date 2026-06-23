"""
Word (DOCX) file parser.
"""

from typing import Dict, Any, Optional
from .base_parser import BaseFileParser


class DocxParser(BaseFileParser):
    """DOCX file parser."""

    def parse(self, file_content: bytes, filename: str, **kwargs) -> Dict[str, Any]:
        """Parse a DOCX file."""
        try:
            from docx import Document
            from io import BytesIO

            docx_file = BytesIO(file_content)
            doc = Document(docx_file)

            # Extract text from all paragraphs.
            paragraphs = [para.text for para in doc.paragraphs if para.text.strip()]
            text = "\n\n".join(paragraphs)

            # Extract text from tables.
            table_texts = []
            for table in doc.tables:
                for row in table.rows:
                    row_texts = [
                        cell.text.strip() for cell in row.cells if cell.text.strip()
                    ]
                    if row_texts:
                        table_texts.append(" | ".join(row_texts))

            if table_texts:
                text += "\n\n[Table content]\n" + "\n".join(table_texts)

            # Extract metadata.
            metadata = self._extract_basic_metadata(filename)
            metadata["file_type"] = (
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            )
            metadata["paragraph_count"] = len(paragraphs)
            metadata["table_count"] = len(doc.tables)
            metadata["character_count"] = len(text)

            # Word document properties.
            core_props = doc.core_properties
            if core_props.title:
                metadata["title"] = core_props.title
            if core_props.author:
                metadata["author"] = core_props.author
            if core_props.subject:
                metadata["subject"] = core_props.subject
            if core_props.created:
                metadata["creation_date"] = core_props.created.isoformat()
            if core_props.modified:
                metadata["modified_date"] = core_props.modified.isoformat()

            return {
                "text": text,
                "metadata": metadata,
                "success": True,
                "error": None,
            }
        except ImportError:
            return {
                "text": "",
                "metadata": self._extract_basic_metadata(filename),
                "success": False,
                "error": "python-docx is not installed; please run: pip install python-docx",
            }
        except Exception as e:
            return {
                "text": "",
                "metadata": self._extract_basic_metadata(filename),
                "success": False,
                "error": f"DOCX parsing failed: {str(e)}",
            }

    def can_parse(self, filename: str, mime_type: Optional[str] = None) -> bool:
        """Determine whether the file is a DOCX file.

        Note: python-docx can only read OOXML-format .docx files; it cannot
        parse the legacy binary .doc format (application/msword). For that
        reason we no longer claim .doc here, to avoid mis-routing the old
        format to this parser and raising a confusing failure.
        """
        if mime_type:
            return (
                mime_type
                == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            )
        return filename.lower().endswith(".docx")
